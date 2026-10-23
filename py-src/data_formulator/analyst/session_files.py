"""The session as files (design-docs/60, 65): one view the agent lists, reads, and searches.

- ``files/<path>``: workspace files, in folders.
- ``session/<path>``: the session's state as files, from the snapshot the
  application sends with the request (``session_files``): charts and reports the
  agent may edit, tables and threads it may read.
- ``scratch/<path>``: execution intermediates and terminal output.
- ``sources/<source_id>/<path>``: connected-source catalogs, one level per listing; a table's file is its
  metadata view (``session_sources``). Searched with everything, after the session, from the cache.
- ``workflows/<path>``: the user's saved workflows, the built-in examples and server workflows (read-only),
  with ``workflows/schedules.json``; listed and searched only when a path names the folder.

Each path is read in the unit that fits it (lines, table rows, PDF pages, Excel sheets); see
``session_fs`` for formats and search.
"""

from __future__ import annotations

import fnmatch
import hashlib
import json
import logging
import re
from dataclasses import dataclass
from itertools import islice
from typing import Any, Iterator

from data_formulator.analyst import session_fs

logger = logging.getLogger(__name__)

MAX_TEXT_BYTES = 2_000_000
READ_LINES = 400
MAX_BATCH_FILES = 20
MAX_BATCH_CHARS = 60_000
MAX_MATCHES = 100
MAX_LISTED = 500
WORKFLOWS = "workflows/"
SOURCES = "sources/"
DEFAULT_ROOTS = ("files", "session", "scratch")
SCHEDULES = "workflows/schedules.json"
_TEXT_SUFFIXES = (
    ".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".py", ".sql", ".r", ".js", ".jsx", ".ts", ".tsx", ".html", ".htm", ".css", ".xml", ".svg", ".log", ".sh",
)


def content_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def helper_module_name(path: str) -> str | None:
    """The module a Python helper file imports as (``files/lib/clean.py`` -> ``lib.clean``), or None."""
    if not path.startswith("files/") or not path.endswith(".py"):
        return None
    parts = path[len("files/"):-len(".py")].split("/")
    if parts[-1] == "__init__":
        parts = parts[:-1]
    return ".".join(parts) if parts and all(part.isidentifier() for part in parts) else None


@dataclass
class Entry:
    path: str
    size: int | None
    editable: bool
    display_name: str | None = None
    note: str | None = None


class SessionFiles:
    """The files of one run's session (``ctx``: workspace + payload)."""

    def __init__(self, ctx: Any):
        self.ctx = ctx
        self.workspace = getattr(ctx, "workspace", None)
        self.payload = getattr(ctx, "payload", None) or {}
        self._workflows: dict[str, dict[str, Any]] | None = None
        self._sources: Any = None
        self._workflow_store: Any = None

    # -- session/ snapshot ---------------------------------------------------

    @property
    def layer(self):
        from data_formulator.analyst.session_layer import session_layer
        return session_layer(self.payload)

    def session_view(self, path: str) -> dict[str, Any] | None:
        layer = self.layer
        return layer.view(path) if layer is not None else None

    # -- sources/ --------------------------------------------------------------

    @property
    def sources(self):
        if self._sources is None:
            from data_formulator.analyst.session_sources import Sources
            self._sources = Sources(self.workspace, self.layer)
        return self._sources

    # -- workflows/ ----------------------------------------------------------

    def workflows_available(self) -> bool:
        if self.payload.get("in_workflow_run") or self.ctx is None:
            return False
        from data_formulator.analyst.skills.configure.automation import workflows_unavailable
        try:
            return workflows_unavailable() is None
        except Exception:  # no application context (tests, offline tools)
            return False

    def workflows(self) -> dict[str, dict[str, Any]]:
        """Workflows by path (``workflows/<store path>``) with their schedules, and the schedules view."""
        if self._workflows is None:
            self._workflows = {}
            if self.workflows_available():
                try:
                    self._workflows = self._load_workflows()
                except Exception:
                    logger.debug("workflows could not be listed", exc_info=True)
        return self._workflows

    def _load_workflows(self) -> dict[str, dict[str, Any]]:
        from data_formulator.analyst.skills.configure import automation
        self._workflow_store = automation.workflow_store(self.ctx)
        try:
            scheduling = automation.list_schedules(self.ctx)
        except Exception as error:  # the workflows are still listed and readable
            logger.debug("schedules could not be listed", exc_info=True)
            scheduling = {"available": False, "error": str(error)[:200]}
        schedules = scheduling.get("schedules") or []
        found = {f"{WORKFLOWS}{item['path']}": {
            **item, "schedules": [schedule for schedule in schedules if schedule["workflow"] == item["path"]]}
            for item in self._workflow_store.list_all()}
        found[SCHEDULES] = {"name": "Workflow schedules", "scheduling": scheduling}
        return found

    def _workflow_detail(self, path: str) -> str:
        item = self.workflows().get(path) or {}
        if path == SCHEDULES:
            scheduling = item.get("scheduling") or {}
            if not scheduling.get("available"):
                return f"scheduling unavailable: {scheduling.get('error', '')}".rstrip(": ")
            count = len(scheduling.get("schedules") or [])
            return f"{count} {'schedule' if count == 1 else 'schedules'}; server models; IDs for propose_schedule"
        if item.get("error"):
            return f"invalid: {item['error'][:120]}"
        parts = [{"user": "saved by the user", "demo": "built-in example", "server": "server workflow"}.get(
            item.get("origin"), "")]
        if item.get("parameters"):
            parts.append("parameters: " + ", ".join(str(parameter.get("name", "")) for parameter in item["parameters"]))
        for schedule in item.get("schedules") or []:
            state = f"next {schedule['next_at']}" if schedule.get("enabled") and schedule.get("next_at") else \
                ("on" if schedule.get("enabled") else "paused")
            parts.append(f'schedule "{schedule["name"]}" ({schedule["id"]}): {schedule["cadence"]}, {state}')
        return " · ".join(part for part in parts if part)

    def _workflow_text(self, path: str) -> str:
        item = self.workflows().get(path)
        if item is None:
            raise FileNotFoundError(path)
        if path == SCHEDULES:
            return json.dumps(item["scheduling"], indent=2, ensure_ascii=False, default=str)
        return self._workflow_store.read(path[len(WORKFLOWS):])

    # -- listing -------------------------------------------------------------

    def entries(self, roots: Any = None) -> Iterator[Entry]:
        """The paths under the top-level folders ``roots`` (default: files/, session/ and scratch/;
        workflows/ only when named)."""
        roots = set(DEFAULT_ROOTS if roots is None else roots)
        if self.workspace is not None and "files" in roots:
            for item in sorted(self.workspace.list_workspace_files(), key=lambda item: item.name.lower()):
                module = helper_module_name(f"files/{item.name}")
                note = None
                if module:
                    note = f"importable as {module}" if getattr(item, "code_signature", None) else \
                        "not importable until you save it (uploaded or edited by the user)"
                yield Entry(f"files/{item.name}", item.file_size,
                            item.origin == "agent" and item.edit_policy == "agent_editable",
                            item.display_name, note)
        layer = self.layer if "session" in roots else None
        for path in layer.paths() if layer is not None else []:
            # Views are built when read; listing names them without building them.
            editable = path.startswith(("session/charts/", "session/reports/"))
            yield Entry(path, None, editable)
        if self.workspace is not None and "scratch" in roots:
            for path in self.workspace.list_scratch_files():
                try:
                    size = self.workspace.resolve_scratch_file(path[len("scratch/"):]).stat().st_size
                except (OSError, ValueError):
                    continue
                yield Entry(path, size, False)
        if "workflows" in roots:
            for path, item in self.workflows().items():
                yield Entry(path, None, False, item.get("name"))

    def list(self, path: str = "") -> dict[str, Any]:
        """Every path under ``path`` (a folder, prefix or glob) with the name the user sees, its kind and size.
        ``workflows/`` is listed only when ``path`` names it."""
        pattern = _normalize(path) if isinstance(path, str) and any(char in path for char in "*?[") else None
        prefix = "" if pattern else _prefix(path)
        if (pattern or prefix).startswith(SOURCES):
            # Catalogs are listed one level at a time; a glob lists matching tables at any depth.
            found = self.sources.list(pattern or prefix) if self.sources.available else []
            return {"path": pattern or prefix, "count": len(found), "files": found[:MAX_LISTED],
                    **({"truncated": True} if len(found) > MAX_LISTED else {})}
        named = (pattern or prefix).startswith(WORKFLOWS[:-1])
        listed = [entry for entry in self.entries(roots_of([pattern or prefix]))
                  if (fnmatch.fnmatch(entry.path, pattern) if pattern else entry.path.startswith(prefix))]
        items = [self._listing(entry) for entry in listed[:MAX_LISTED]]
        result = {"path": pattern or prefix or "", "count": len(listed), "files": items,
                  **({"truncated": True} if len(listed) > MAX_LISTED else {})}
        if not named and not prefix and not pattern:
            also = []
            if self.sources.available:
                also.append("sources/: connected-source catalogs, one level per listing (list_files path sources/).")
            if self.workflows_available():
                also.append("workflows/: the user's saved and built-in workflows and their schedules, listed "
                            "only when named (list_files path workflows/).")
            if also:
                result["also"] = " ".join(also)
        return result

    def _listing(self, entry: Entry) -> dict[str, Any]:
        kind = session_fs.kind_of(entry.path)
        name = entry.display_name or self.display_name(entry.path)
        item: dict[str, Any] = {"path": entry.path, **({"name": name} if name else {}), "kind": kind}
        try:
            detail = self._detail(entry.path, kind)
        except Exception:
            detail = ""
        if detail:
            item["detail"] = detail
        if entry.size is not None:
            item["size"] = entry.size
        if entry.editable:
            item["editable"] = True
        if entry.note:
            item["note"] = entry.note
        return item

    def _detail(self, path: str, kind: str) -> str:
        """Rows, pages or sheets: what tells the agent which read_file parameters apply."""
        if path.startswith(WORKFLOWS):
            return self._workflow_detail(path)
        if kind == "table" and self.layer is not None:
            rows, columns = self.layer.table_shape(path[len("session/tables/"):-len(".json")])
            return f"{rows:,} rows × {columns} columns" if rows is not None else ""
        if not path.startswith("files/") or self.workspace is None:
            return ""
        name = path[len("files/"):]
        if kind == "PDF":
            return f"{len(session_fs.pdf_reader(self.workspace.read_workspace_file(name)[1]).pages)} pages"
        if kind == "Excel":
            from openpyxl import load_workbook
            import io
            workbook = load_workbook(io.BytesIO(self.workspace.read_workspace_file(name)[1]), read_only=True)
            try:
                return "sheets: " + ", ".join(f"{sheet.title} ({max((sheet.max_row or 1) - 1, 0):,} rows)"
                                              for sheet in workbook.worksheets)
            finally:
                workbook.close()
        if kind == "data file":
            import pyarrow.parquet as pq
            local = self.workspace.local_workspace_file(name) if hasattr(self.workspace, "local_workspace_file") else None
            if local is None:
                import io
                local = io.BytesIO(self.workspace.read_workspace_file(name)[1])
            metadata = pq.ParquetFile(local).metadata
            return f"{metadata.num_rows:,} rows × {metadata.num_columns} columns"
        return ""

    def display_name(self, path: str) -> str | None:
        """The name the user sees for ``path``."""
        if path.startswith("session/"):
            layer = self.layer
            try:
                return layer.name_of(path) if layer is not None else None
            except Exception:
                return None
        if path.startswith("files/") and self.workspace is not None:
            item = self.workspace.get_metadata().files.get(path[len("files/"):])
            return item.display_name if item is not None else None
        if path.startswith(WORKFLOWS):
            return (self.workflows().get(path) or {}).get("name")
        if path.startswith(SOURCES) and self.sources.available:
            table = self.sources.table(path)
            return table.name if table is not None and table.name != table.segments[-1] else None
        return None

    def iter_lines(self, path: str, max_bytes: int | None = None) -> Iterator[tuple[int, str]]:
        """(number, line) of a text path, streamed from disk where it can be, without the size limit."""
        path = _normalize(path)
        if session_fs.kind_of(path) == "Word":
            yield from enumerate(self.text(path).splitlines(), start=1)
            return
        local = self._local_path(path)
        if local is None:
            text = self.text(path) if not path.startswith(("files/", "scratch/")) else \
                self._raw(path).decode("utf-8", errors="replace")
            for number, line in enumerate(text.splitlines(), start=1):
                yield number, line
            return
        read = 0
        with open(local, "r", encoding="utf-8", errors="replace") as handle:
            for number, line in enumerate(handle, start=1):
                read += len(line)
                yield number, line.rstrip("\r\n")
                if max_bytes is not None and read >= max_bytes:
                    return

    def _local_path(self, path: str):
        if self.workspace is None:
            return None
        if path.startswith("files/") and hasattr(self.workspace, "local_workspace_file"):
            return self.workspace.local_workspace_file(path[len("files/"):])
        if path.startswith("scratch/"):
            try:
                local = self.workspace.resolve_scratch_file(path[len("scratch/"):])
            except (OSError, ValueError):
                return None
            return local if local.is_file() else None
        return None

    def _raw(self, path: str) -> bytes:
        if path.startswith("files/"):
            try:
                return self.workspace.read_workspace_file(path[len("files/"):])[1]
            except (FileNotFoundError, KeyError):
                raise FileNotFoundError(path) from None
        try:
            return self.workspace.resolve_scratch_file(path[len("scratch/"):]).read_bytes()
        except (OSError, ValueError):
            raise FileNotFoundError(path) from None

    def _size(self, path: str) -> int | None:
        local = self._local_path(path)
        if local is not None:
            return local.stat().st_size
        if path.startswith("files/") and self.workspace is not None:
            item = self.workspace.get_metadata().files.get(path[len("files/"):])
            return item.file_size if item is not None else None
        return None

    # -- reading -------------------------------------------------------------

    def text(self, path: str) -> str:
        """A file's full text; raises ``FileNotFoundError`` or ``ValueError`` (not text)."""
        path = _normalize(path)
        if path.startswith("session/"):
            item = self.session_view(path)
            if item is None:
                raise FileNotFoundError(path)
            return str(item.get("content", ""))
        if path.startswith(WORKFLOWS):
            return self._workflow_text(path)
        if path.startswith(SOURCES):
            if not self.sources.available:
                raise FileNotFoundError(path)
            return self.sources.view(path)
        if self.workspace is None:
            raise FileNotFoundError(path)
        if path.startswith("files/"):
            try:
                metadata, content = self.workspace.read_workspace_file(path[len("files/"):])
            except (FileNotFoundError, KeyError):
                raise FileNotFoundError(path) from None
            if session_fs.kind_of(path) == "Word":
                from data_formulator.datalake.workspace_file_content import extract_workspace_file_text
                try:
                    return extract_workspace_file_text(metadata.name, content, metadata.media_type).content
                except Exception as error:
                    raise ValueError(f"{path} could not be read as a document: {error}") from None
        elif path.startswith("scratch/"):
            try:
                content = self.workspace.resolve_scratch_file(path[len("scratch/"):]).read_bytes()
            except (OSError, ValueError):
                raise FileNotFoundError(path) from None
        else:
            raise FileNotFoundError(path)
        return _decode(path, content)

    def read(self, path: str, offset: int = 1, limit: int = READ_LINES) -> dict[str, Any]:
        """Lines ``offset`` .. ``offset + limit - 1`` of a text path; a large file is paged, never loaded whole."""
        path = _normalize(path)
        if not isinstance(offset, int) or offset < 1:
            raise ValueError("offset is the first line to read, starting at 1")
        if not isinstance(limit, int) or not 1 <= limit <= 2000:
            raise ValueError("limit must be between 1 and 2000 lines")
        size = self._size(path) if path.startswith(("files/", "scratch/")) and session_fs.kind_of(path) != "Word" else None
        if size is not None and size > MAX_TEXT_BYTES:
            return self._read_large(path, offset, limit, size)
        try:
            text = self.text(path)
        except FileNotFoundError:
            raise ValueError(f"{path} does not exist. Use list_files to see the session's files.") from None
        lines = text.splitlines(keepends=True)
        chunk = lines[offset - 1:offset - 1 + limit]
        editable = self._editable(path)
        result: dict[str, Any] = {
            "path": path, "content": "".join(chunk), "start_line": offset,
            "end_line": offset - 1 + len(chunk), "total_lines": len(lines),
            "content_hash": content_hash(text) if path.startswith("session/") else self._file_hash(path, text),
            **({"editable": True} if editable else {}),
        }
        if offset - 1 + len(chunk) < len(lines):
            result["more"] = f"Read on with offset={offset + len(chunk)}."
        note = (self.session_view(path) or {}).get("note") if path.startswith("session/") else None
        if path.startswith(WORKFLOWS) and path != SCHEDULES:
            origin = (self.workflows().get(path) or {}).get("origin")
            note = ("Read-only. To revise it, propose the complete revised definition with propose_workflow"
                    + (f' and replaces="{path}".' if origin == "user" else
                       "; built-in and server workflows are not replaced, so the user saves the revision as a new workflow."))
        if note:
            result["note"] = note
        return result

    def _read_large(self, path: str, offset: int, limit: int, size: int) -> dict[str, Any]:
        chunk = [line for _, line in islice(self.iter_lines(path), offset - 1, offset - 1 + limit)]
        total = sum(1 for _ in self.iter_lines(path))
        end = offset - 1 + len(chunk)
        result: dict[str, Any] = {
            "path": path, "content": "\n".join(chunk) + ("\n" if chunk else ""), "start_line": offset,
            "end_line": end, "total_lines": total,
            "more": (f"PARTIAL view of a large file ({size:,} bytes, {total:,} lines): lines {offset}-{end}. "
                     f"Read on with offset={end + 1}, search it with search, or process it with Python."
                     if end < total else f"Large file ({size:,} bytes); this is its end."),
        }
        if path.startswith("files/") and self.workspace is not None:
            item = self.workspace.get_metadata().files.get(path[len("files/"):])
            if item is not None:
                result["content_hash"] = item.content_hash
        return result

    def read_items(self, items: Any) -> dict[str, Any]:
        """Several items at once, each ``{path, ...unit parameters}``; globs expand with the same parameters."""
        if isinstance(items, (str, dict)):
            items = [items]
        if not isinstance(items, list) or not items:
            raise ValueError('files is a list of items such as [{"path": "session/tables/sales.json", "rows": "1-50"}, '
                             '{"path": "files/report.pdf", "pages": "3-5"}]')
        wanted: list[dict[str, Any]] = []
        unmatched = []
        for item in items:
            item = {"path": item} if isinstance(item, str) else item
            if not isinstance(item, dict) or not isinstance(item.get("path"), str) or not item["path"].strip():
                raise ValueError('Each item needs a path, e.g. {"path": "files/notes/plan.md"}')
            pattern = item["path"].strip()
            if any(char in pattern for char in "*?["):
                glob = _normalize(pattern)
                found = [entry.path for entry in self.entries() if fnmatch.fnmatch(entry.path, glob)]
                unmatched += [] if found else [pattern]
                wanted += [{**item, "path": path} for path in found]
            else:
                wanted.append(item)
        seen: set[str] = set()
        results: list[dict[str, Any]] = []
        used = 0
        for index, item in enumerate(wanted):
            key = json.dumps(item, sort_keys=True, default=str)
            if key in seen:
                continue
            seen.add(key)
            if index >= MAX_BATCH_FILES:
                results.append({"path": item["path"], "not_read": f"more than {MAX_BATCH_FILES} items; read it separately"})
                continue
            try:
                result = self.read_item(item)
            except ValueError as error:
                results.append({"path": item["path"], "error": str(error)})
                continue
            size = len(str(result.get("content", "")))
            if used + size > MAX_BATCH_CHARS and results:
                results.append({"path": item["path"], "not_read": "over this call's size budget; read it separately"})
                continue
            used += size
            results.append(result)
        response: dict[str, Any] = {"files": results}
        if unmatched:
            response["no_match"] = unmatched
        return response

    def read_item(self, item: dict[str, Any]) -> dict[str, Any]:
        """One path, read in the unit that fits it; parameters that do not fit fall back to defaults with a warning."""
        raw_path = item["path"].strip()
        sheet = item.get("sheet")
        if "#" in raw_path and session_fs.kind_of(raw_path.split("#", 1)[0]) == "Excel":
            raw_path, sheet = raw_path.split("#", 1)
        path = _normalize(raw_path)
        kind = session_fs.kind_of(path)
        if kind in ("image", "binary"):
            raise ValueError(f"{path} is a {kind} file read_file cannot show; process it with execute_python_script "
                             f"(for example PIL for images, zipfile for archives, python-docx for Word files).")
        given = {key: item[key] for key in session_fs.UNIT_PARAMS if item.get(key) is not None}
        if sheet is not None:
            given["sheet"] = sheet
        allowed = session_fs.PARAMS_BY_KIND.get(kind, ("offset", "limit"))
        ignored = [key for key in given if key not in allowed]
        ignored += [key for key in item if key != "path" and key not in session_fs.UNIT_PARAMS]
        params = {key: value for key, value in given.items() if key in allowed}
        if kind in ("table", "Excel", "data file") and "columns" in params and "rows" not in params:
            params["rows"] = "1-50"
        name = self.display_name(path)
        if kind in ("table", "Excel", "data file", "CSV") and "rows" in params:
            source = session_fs.table_source(self, path, sheet=params.get("sheet"))
            if source is None:
                raise ValueError(f"{path} does not exist. Use list_files to see the session's files.")
            result = {"path": path, **source.read_rows(params["rows"], params.get("columns"))}
            if params.get("sheet"):
                result["sheet"] = params["sheet"]
        elif kind in ("Excel", "data file"):
            result = {"path": path, **self._data_file_view(path, kind)}
        elif kind == "PDF":
            if self.workspace is None:
                raise ValueError(f"{path} does not exist.")
            try:
                content = self.workspace.read_workspace_file(path[len("files/"):])[1]
            except (FileNotFoundError, KeyError):
                raise ValueError(f"{path} does not exist. Use list_files to see the session's files.") from None
            result = {"path": path, **session_fs.read_pdf(content, params.get("pages"))}
        else:
            result = self.read(path, params.get("offset", 1), params.get("limit", READ_LINES))
        if name and "name" not in result:
            result = {"path": result.pop("path"), "name": name, **result}
        if ignored:
            fits = ", ".join(allowed) if allowed else "no unit parameters"
            unit = {"table": "a table", "Excel": "an Excel file", "data file": "a data file", "CSV": "a CSV file",
                    "PDF": "a PDF"}.get(kind, "a text file, read by lines")
            result["warning"] = f"Ignored {', '.join(ignored)}: this is {unit} ({fits}); read with the defaults."
        return result

    def _data_file_view(self, path: str, kind: str) -> dict[str, Any]:
        """An Excel or parquet file's shape: sheets or columns, and a few sample rows."""
        name = path[len("files/"):]
        if self.workspace is None or self.workspace.get_metadata().files.get(name) is None:
            raise ValueError(f"{path} does not exist. Use list_files to see the session's files.")
        if kind == "Excel":
            sheets = session_fs.excel_sheets(self.workspace.read_workspace_file(name)[1])
            parts = []
            for sheet in sheets:
                source = session_fs.table_source(self, path, sheet=sheet["sheet"])
                sample = source.sample(3)["content"] if source is not None and sheet["rows"] else ""
                parts.append(f'Sheet "{sheet["sheet"]}" — {sheet["rows"]:,} rows: {", ".join(sheet["columns"])}\n{sample}')
            return {"content": "\n\n".join(parts), "sheets": [sheet["sheet"] for sheet in sheets],
                    "note": f'Read a sheet\'s rows with sheet and rows, e.g. {{"path": "{path}", "sheet": '
                            f'"{sheets[0]["sheet"] if sheets else "Sheet1"}", "rows": "1-50"}}.'}
        source = session_fs.table_source(self, path)
        columns = source.columns()
        total = source.row_count()
        return {"content": f"{total:,} rows × {len(columns)} columns: {', '.join(columns)}\n"
                           + source.sample(5)["content"],
                "note": f'Read rows with rows, e.g. {{"path": "{path}", "rows": "1-50"}}; Python reads {path}.'}

    def _file_hash(self, path: str, text: str) -> str:
        if path.startswith("files/") and self.workspace is not None:
            item = self.workspace.get_metadata().files.get(path[len("files/"):])
            if item is not None:
                return item.content_hash
        return content_hash(text)

    def _editable(self, path: str) -> bool:
        if path.startswith("session/"):
            return bool((self.session_view(path) or {}).get("editable"))
        if path.startswith("files/") and self.workspace is not None:
            item = self.workspace.get_metadata().files.get(path[len("files/"):])
            return bool(item and item.origin == "agent" and item.edit_policy == "agent_editable")
        return False

    # -- searching -----------------------------------------------------------

    def search(self, query: str, path: Any = None, *, regex: bool = False, max_results: int = 20) -> str:
        """Search tables (names, columns, every row), files and session items; see ``session_fs.search``."""
        return session_fs.search(self, query, path, regex=regex, max_results=max_results)


def _normalize(path: Any) -> str:
    if not isinstance(path, str) or not path.strip():
        raise ValueError("path is required, e.g. files/notes/plan.md or session/charts/<id>.json")
    path = path.strip().lstrip("/")
    if path.split("/")[0] not in ("files", "session", "scratch", "sources", "workflows") or ".." in path.split("/"):
        raise ValueError(f"{path} is not a session file; paths start with files/, session/, scratch/, sources/, "
                         "or workflows/")
    return path


def roots_of(paths: list[Any]) -> list[str] | None:
    """The top-level folders a set of paths or globs can reach; None for the default folders."""
    roots: list[str] = []
    for path in paths:
        first = str(path or "").strip().lstrip("/").split("/", 1)[0]
        if not first or any(char in first for char in "*?["):
            return None
        roots.append(first)
    return roots or None


def _prefix(path: Any) -> str:
    if path in (None, "", "/"):
        return ""
    prefix = _normalize(path)
    return prefix if prefix.endswith("/") or "." in prefix.rsplit("/", 1)[-1] else prefix + "/"


def _decode(path: str, content: bytes) -> str:
    if len(content) > MAX_TEXT_BYTES:
        raise ValueError(f"{path} is larger than {MAX_TEXT_BYTES:,} bytes; read it with Python")
    if b"\x00" in content and not path.lower().endswith(_TEXT_SUFFIXES):
        raise ValueError(f"{path} is not a text file; read it with Python or the data tools")
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        raise ValueError(f"{path} is not UTF-8 text; read it with Python or the data tools") from None


NOTES_FOLDER = "notes/"
MAX_NOTE_CHARS = 4000
MAX_NOTES_CHARS = 12000


def render_session_notes(workspace: Any) -> str:
    """The ``[SESSION NOTES]`` block: the agent's notes in ``files/notes/``, so decisions carry across turns."""
    if workspace is None:
        return ""
    try:
        notes = sorted((item for item in workspace.list_workspace_files() if item.name.startswith(NOTES_FOLDER)),
                       key=lambda item: item.name)
    except Exception:
        return ""
    if not notes:
        return ""
    lines = ["[SESSION NOTES]",
             "Your working notes for this session (files/notes/); keep them current with edit_file."]
    used = 0
    for item in notes:
        try:
            _, content = workspace.read_workspace_file(item.name)
            text = _decode(f"files/{item.name}", content)
        except Exception:
            lines.append(f"\n## files/{item.name} (not text)")
            continue
        room = min(MAX_NOTE_CHARS, MAX_NOTES_CHARS - used)
        if room <= 0:
            lines.append(f"\n## files/{item.name} (not shown; read_file it)")
            continue
        shown = text[:room]
        used += len(shown)
        more = f"\n… ({len(text) - room:,} more characters; read_file it)" if len(text) > room else ""
        editable = item.origin == "agent" and item.edit_policy == "agent_editable"
        state = (f"content_hash {item.content_hash}" if editable
                 else "written by the user; read-only to you, so add yours in a new notes file")
        lines.append(f"\n## files/{item.name} ({state})\n{shown.rstrip()}{more}")
    return "\n".join(lines)


MAX_INLINE_CHARS = 8000
MAX_THREAD_CHARS = 6000
MAX_TABLE_CHARS = 5000
MAX_TABLES_CHARS = 10000


def _inline(label: str, path: str, files: "SessionFiles", limit: int, tail: bool = False) -> str | None:
    try:
        read = files.read(path, 1, 2000)
    except ValueError:
        return None
    text = read["content"]
    meta = f"content_hash {read['content_hash']}" + (", you can edit it" if read.get("editable") else "")
    if read.get("note"):
        meta += f". {read['note']}"
    if len(text) > limit:
        if tail:
            head, _, body = text.partition("\n")
            text = head + "\n… (earlier steps: read_file " + path + ")\n" + body[-(limit - len(head)):].split("\n", 1)[-1]
        else:
            text = text[:limit] + f"\n… ({len(text) - limit:,} more characters: read_file {path})"
    return f"{label} `{path}` ({meta})\n{text.rstrip()}"


def render_turn_context(ctx: Any, history: list[dict[str, Any]]) -> str:
    """What a continued request adds to the conversation (design-docs/63): what changed since the agent last
    looked and where the user is now; what is on the canvas when it is new or changed, with the data behind it
    the conversation has not seen; the thread's latest steps when the user moved to another thread; and the
    notes when they changed."""
    from data_formulator.analyst.session_layer import session_layer
    from data_formulator.analyst.session_layer.layer import seen_versions

    layer = session_layer(getattr(ctx, "payload", None))
    if layer is None:
        return ""
    seen = seen_versions(history)
    parts = [layer.changes(seen, always_focus=True)]
    current = layer.versions
    path, thread = layer.focus()
    files = SessionFiles(ctx)
    said = json.dumps(history, ensure_ascii=False, default=str)
    if path and (seen.get("@focus") != path or seen.get(path) != current.get(path)):
        parts.append(_inline("[ON THE CANVAS]", path, files, MAX_INLINE_CHARS))
        room = MAX_TABLES_CHARS
        for table in layer.tables_behind(path):
            # Tables the conversation already has (the agent made or read them) are not repeated.
            name = table.rsplit("/", 1)[-1][:-len(".json")]
            if room <= 0 or re.search(rf"(?<![\w-]){re.escape(name)}(?![\w-])", said):
                continue
            block = _inline("[ITS DATA]", table, files, min(MAX_TABLE_CHARS, room))
            if block:
                parts.append(block)
                room -= len(block)
    thread_path = f"session/threads/{thread}.md" if thread else None
    focus_lines = [line for line in said.replace("\\n", "\n").split("\n") if line.startswith("Focus: ")]
    if thread_path and (not focus_lines or thread_path not in focus_lines[-1]):
        parts.append(_inline("[THE THREAD YOU ARE CONTINUING]", thread_path, files, MAX_THREAD_CHARS, tail=True))
    if any(name.startswith("files/" + NOTES_FOLDER) and current.get(name) != seen.get(name) for name in current):
        parts.append(render_session_notes(getattr(ctx, "workspace", None)))
    return "\n\n".join(part for part in parts if part)


def render_session_context(ctx: Any, *, focus: bool = True) -> str:
    """What a request shows of the session (design-docs/61): the index, the notes, what is on the canvas,
    and the latest steps of the thread being continued, each as the file it is. Runs not started from the
    canvas (workflows) leave out the focus."""
    from data_formulator.analyst.session_layer import session_layer

    layer = session_layer(getattr(ctx, "payload", None))
    workspace = getattr(ctx, "workspace", None)
    parts = []
    if layer is not None:
        parts.append("[SESSION INDEX]\n" + layer.index().rstrip())
    notes = render_session_notes(workspace)
    if notes:
        parts.append(notes)
    if layer is not None and focus:
        files = SessionFiles(ctx)
        path, thread = layer.focus()
        if path:
            parts.append(_inline("[ON THE CANVAS]", path, files, MAX_INLINE_CHARS))
            # The data behind a chart or table on the canvas, which most follow-ups read first.
            room = MAX_TABLES_CHARS
            for table in layer.tables_behind(path):
                if room <= 0:
                    break
                block = _inline("[ITS DATA]", table, files, min(MAX_TABLE_CHARS, room))
                if block:
                    parts.append(block)
                    room -= len(block)
        if thread:
            parts.append(_inline("[THE THREAD YOU ARE CONTINUING]", f"session/threads/{thread}.md", files,
                                 MAX_THREAD_CHARS, tail=True))
    return "\n\n".join(part for part in parts if part)
